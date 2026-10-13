"""Editable PowerPoint version of the 2-slide ContextOps pitch (mirrors project/slides/*.html).

Usage: python build_pptx.py  ->  writes ContextOps-pitch.pptx next to this script.
Layout is authored on a 1920x1080 px canvas and converted to EMU.
"""
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Pt

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "ContextOps-pitch.pptx")
HEAD, BODY = "Space Grotesk", "IBM Plex Sans"
NAVY, MINT, PAPER = "0F1B2D", "4FD1A1", "F6F5F0"


def px(v):  # 1920 px = 12,192,000 EMU
    return Emu(int(round(v * 6350)))


def rgb(h):
    return RGBColor.from_string(h)


def notes_from(slide_id):
    html = open(os.path.join(ROOT, "project", "slides", f"{slide_id}.html"), encoding="utf-8").read()
    m = re.search(r"<aside>(.*?)</aside>", html, re.S)
    return m.group(1).strip() if m else ""


def text(slide, x, y, w, h, runs, size, color, font=BODY, bold=False, spacing=None, line=1.15):
    """runs: a string or a list of (text, bold) tuples, rendered as one paragraph."""
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.line_spacing = line
    for t, b in ([(runs, bold)] if isinstance(runs, str) else runs):
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size * 0.5)  # px -> pt on a 1920-wide canvas
        r.font.name = font
        r.font.bold = b
        r.font.color.rgb = rgb(color)
        if spacing is not None:
            r.font._element.set("spc", str(spacing))  # _element is the <a:rPr> node
    return box


def bar(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    return shp


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]

# Slide 1: the problem and why it matters
s1 = prs.slides.add_slide(blank)
background(s1, NAVY)
text(s1, 128, 128, 1664, 40, "CONTEXTOPS · THE PROBLEM", 26, MINT, bold=True, spacing=500)
text(s1, 128, 190, 1664, 190,
     "Companies are pouring money into AI coding tools, and can’t tell what it buys.",
     76, PAPER, font=HEAD, bold=True, line=1.05)
stats = [("86%", "of engineering leaders can’t say which AI tools pay off"),
         ("2028", "AI coding costs pass the average developer salary"),
         ("98%", "of FinOps teams now manage AI spend, up from 31%")]
for i, (big, label) in enumerate(stats):
    x = 128 + i * 576
    text(s1, x, 500, 512, 110, big, 96, MINT, font=HEAD, bold=True, line=1.0)
    text(s1, x, 614, 500, 90, label, 28, "C3CCD6", line=1.2)
bar(s1, 128, 820, 4, 100, MINT)
text(s1, 164, 828, 1628, 100,
     "Why it matters: engineering managers now own a budget that grows every month, "
     "with no way to tie it to what actually shipped.", 34, PAPER, line=1.2)
text(s1, 128, 990, 1664, 32,
     "Sources: DX survey of 275 leaders · Gartner, Jun 2026 · FinOps Foundation, State of FinOps 2026",
     24, "8A97A8")
s1.notes_slide.notes_text_frame.text = notes_from("cover")

# Slide 2: progress this semester
s2 = prs.slides.add_slide(blank)
background(s2, PAPER)
text(s2, 128, 128, 1664, 40, "CONTEXTOPS · THIS SEMESTER", 26, "16784F", bold=True, spacing=500)
text(s2, 128, 186, 1664, 90, "From research to a working product", 72, NAVY, font=HEAD, bold=True, line=1.05)
steps = [("01  Researched", "2,200+ engineer posts and 35 startup ideas scored. This one won every round."),
         ("02  Built", "A working MVP that ties Claude Code, Cursor and Copilot spend to merged code, "
                       "and fixes waste in one click."),
         ("03  Next", "20 customer interviews and 5 design partners.")]
y = 330
for title, body in steps:
    text(s2, 128, y, 620, 50, title, 36, "16784F", font=HEAD, bold=True)
    text(s2, 128, y + 54, 620, 120, body, 28, "3A4556", line=1.2)
    y += 200
s2.shapes.add_picture(os.path.join(ROOT, "assets", "mvp.png"), px(820), px(330), px(972), px(471))
s2.notes_slide.notes_text_frame.text = notes_from("traction")

prs.save(OUT)
print("saved", OUT)
